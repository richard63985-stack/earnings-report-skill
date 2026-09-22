# -*- coding: utf-8 -*-
"""Generate Monday-style visual evidence workpapers from a JSON config.

Usage:
    python create_visual_workpaper.py --config path/to/workpaper_config.json

The config is intentionally data-heavy: future companies should change the
sentence/source mapping JSON, not this script.
"""

from __future__ import annotations

import argparse
import os
import re
import hashlib
import json
import math
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break
from PIL import Image, ImageDraw, ImageFont
from playwright.sync_api import sync_playwright



def resolve_path(value: str | Path, base_dir: Path) -> Path:
    path = Path(value)
    if path.is_absolute():
        return path
    return (base_dir / path).resolve()


def has_path_part(path: Path, part: str) -> bool:
    return any(piece.lower() == part.lower() for piece in path.parts)


def validate_output_paths(output_xlsx: Path, process_dir: Path, cfg: dict) -> None:
    if cfg.get("allow_submit_materials_output"):
        return
    blocked = [path for path in (output_xlsx, process_dir) if has_path_part(path, "提交材料")]
    if blocked:
        paths = ", ".join(str(path) for path in blocked)
        raise ValueError(f"Output/process paths must not point into 提交材料 unless allow_submit_materials_output=true: {paths}")


def load_config(path: Path) -> dict:
    with path.open("r", encoding="utf-8-sig") as f:
        cfg = json.load(f)
    cfg["_config_dir"] = str(path.parent.resolve())
    return cfg


def font(size: int) -> ImageFont.ImageFont:
    candidates = [os.environ.get("REPORT_FONT", ""),
        "C:/Windows/Fonts/msyh.ttc", "C:/Windows/Fonts/simsun.ttc",
        "/System/Library/Fonts/PingFang.ttc", "/System/Library/Fonts/STHeiti Medium.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"]
    for candidate in candidates:
        if candidate and Path(candidate).is_file():
            return ImageFont.truetype(candidate, size)
    raise RuntimeError("Chinese font missing; set REPORT_FONT to a CJK .ttf/.ttc/.otf file")

BODY_FONT = font(28)
BOLD_FONT = font(30)
TITLE_FONT = font(34)


def text_width(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.ImageFont) -> int:
    bbox = draw.textbbox((0, 0), text, font=fnt)
    return bbox[2] - bbox[0]


def text_height(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.ImageFont) -> int:
    bbox = draw.textbbox((0, 0), text or "测", font=fnt)
    return bbox[3] - bbox[1]


def is_cjk(ch: str) -> bool:
    return "\u4e00" <= ch <= "\u9fff" or "\u3000" <= ch <= "\u303f" or "\uff00" <= ch <= "\uffef"


def iter_units(text: str) -> list[tuple[str, int, int]]:
    """Return drawable units with original start/end offsets."""
    units: list[tuple[str, int, int]] = []
    i = 0
    while i < len(text):
        ch = text[i]
        if ch == "\n":
            units.append(("\n", i, i + 1))
            i += 1
        elif ch.isspace():
            j = i + 1
            while j < len(text) and text[j].isspace() and text[j] != "\n":
                j += 1
            units.append((" ", i, j))
            i = j
        elif is_cjk(ch):
            units.append((ch, i, i + 1))
            i += 1
        else:
            j = i + 1
            while j < len(text) and (not text[j].isspace()) and (not is_cjk(text[j])):
                j += 1
            units.append((text[i:j], i, j))
            i = j
    return units


def compact_text_with_map(text: str) -> tuple[str, list[int]]:
    compact: list[str] = []
    mapping: list[int] = []
    for idx, ch in enumerate(text):
        if ch.isspace() or ch == "\u00a0":
            continue
        compact.append(ch.lower())
        mapping.append(idx)
    return "".join(compact), mapping


def highlight_ranges(
    text: str,
    highlights: list[str],
    strict: bool = False,
    context: str = "",
) -> list[tuple[int, int]]:
    ranges: list[tuple[int, int]] = []
    compact_body, body_map = compact_text_with_map(text)
    for phrase in highlights:
        if not phrase:
            continue
        compact_phrase, _ = compact_text_with_map(phrase)
        start = compact_body.find(compact_phrase)
        if start >= 0 and compact_phrase:
            end = start + len(compact_phrase) - 1
            ranges.append((body_map[start], body_map[end] + 1))
        elif strict:
            label = f" in {context}" if context else ""
            raise ValueError(f"Highlight not found{label}: {phrase}")
    return ranges


def intersects(start: int, end: int, ranges: list[tuple[int, int]]) -> bool:
    return any(start < r_end and end > r_start for r_start, r_end in ranges)


def wrap_units(
    draw: ImageDraw.ImageDraw,
    text: str,
    fnt: ImageFont.ImageFont,
    max_width: int,
) -> list[list[tuple[str, int, int]]]:
    lines: list[list[tuple[str, int, int]]] = []
    current: list[tuple[str, int, int]] = []
    width = 0
    for unit, start, end in iter_units(text):
        if unit == "\n":
            lines.append(current)
            current = []
            width = 0
            continue
        token_width = text_width(draw, unit, fnt)
        if current and width + token_width > max_width:
            lines.append(current)
            current = []
            width = 0
            if unit == " ":
                continue
        current.append((unit, start, end))
        width += token_width
    if current:
        lines.append(current)
    return lines


def make_text_image(
    sheet: str,
    heading: str,
    body: str,
    out: Path,
    highlights: list[str] | None = None,
    include_sheet_title: bool = True,
    width: int = 900,
) -> dict:
    margin_x = 30
    margin_y = 24
    max_text_width = width - margin_x * 2
    ranges = highlight_ranges(
        body,
        highlights or [],
        strict=bool(highlights),
        context=f"{sheet}/{heading}",
    )

    probe = Image.new("RGB", (width, 200), "white")
    draw = ImageDraw.Draw(probe)
    lines = wrap_units(draw, body, BODY_FONT, max_text_width)
    line_h = text_height(draw, "测", BODY_FONT) + 11
    title_h = text_height(draw, sheet, TITLE_FONT) + 24 if include_sheet_title else 0
    heading_h = text_height(draw, heading, BOLD_FONT) + 12
    height = margin_y * 2 + title_h + heading_h + max(1, len(lines)) * line_h + 12

    img = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(img)
    y = margin_y
    purple = (104, 73, 167)

    if include_sheet_title:
        draw.text((margin_x, y), sheet, font=TITLE_FONT, fill=purple)
        y += text_height(draw, sheet, TITLE_FONT) + 8
        draw.line((margin_x, y, width - margin_x, y), fill=purple, width=3)
        y += 16

    draw.text((margin_x, y), heading, font=BOLD_FONT, fill=(20, 20, 20))
    y += heading_h

    for line in lines:
        x = margin_x
        for unit, start, end in line:
            unit_width = text_width(draw, unit, BODY_FONT)
            if intersects(start, end, ranges):
                draw.rectangle(
                    (x - 1, y - 2, x + unit_width + 1, y + line_h - 3),
                    fill=(216, 216, 216),
                )
            draw.text((x, y), unit, font=BODY_FONT, fill=(40, 40, 40))
            x += unit_width
        y += line_h

    out.parent.mkdir(parents=True, exist_ok=True)
    img.save(out)
    return {"path": str(out), "width": width, "height": height}


def extract_docx_paragraphs(docx_path: Path) -> list[str]:
    ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    with zipfile.ZipFile(docx_path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    paragraphs: list[str] = []
    for par in root.iter(ns + "p"):
        text = "".join(node.text or "" for node in par.iter(ns + "t")).strip()
        if text:
            paragraphs.append(text)
    return paragraphs


def select_body(paragraphs: list[str], section: dict) -> str:
    if section.get("body"):
        if section["body"] not in paragraphs:
            raise ValueError("Section body differs from final Word: " + section.get("sheet", ""))
        return section["body"]
    selector = section.get("body_selector", {})
    prefix = selector.get("prefix")
    contains = selector.get("contains")
    matches = [text for text in paragraphs if (prefix and text.startswith(prefix)) or (contains and contains in text)]
    if len(matches) == 1:
        return matches[0]
    if len(matches) > 1:
        raise ValueError("Body selector is ambiguous; provide the exact Word paragraph")
    raise ValueError(f"Cannot locate body for sheet {section.get('sheet')!r}")


def highlighted_html(source_path: Path, out_path: Path, terms: list[str]) -> Path:
    html = source_path.read_text(encoding="utf-8", errors="ignore")
    soup = BeautifulSoup(html, "html.parser")
    for node in soup(["script", "iframe", "object", "embed"]):
        node.decompose()
    for node in soup.find_all(True):
        for key in list(node.attrs):
            if key.lower().startswith("on"):
                del node.attrs[key]
    if soup.head is None:
        head = soup.new_tag("head")
        if soup.html:
            soup.html.insert(0, head)
        else:
            soup.insert(0, head)
    style = soup.new_tag("style")
    style.string = """
      mark.codex-highlight {
        background: rgba(255, 238, 88, 0.78) !important;
        outline: 4px solid rgba(239, 68, 68, 0.92) !important;
        outline-offset: 2px !important;
        color: inherit !important;
      }
      body { background: #ffffff !important; }
    """
    soup.head.append(style)
    script = soup.new_tag("script")
    script.string = f"""
    (() => {{
      const terms = {json.dumps(terms, ensure_ascii=False)};
      const compactWithMap = (s) => {{
        const compact = [];
        const map = [];
        const raw = String(s || '');
        for (let i = 0; i < raw.length; i++) {{
          const ch = raw[i];
          if (/\\s/.test(ch) || ch === '\\u00a0') continue;
          compact.push(ch.toLowerCase());
          map.push(i);
        }}
        return {{ compact: compact.join(''), map }};
      }};
      let markId = 0;
      const termHits = Object.fromEntries(terms.map((term) => [term, 0]));
      const walker = document.createTreeWalker(document.body || document.documentElement, NodeFilter.SHOW_TEXT);
      const nodes = [];
      while (walker.nextNode()) nodes.push(walker.currentNode);
      const findNext = (raw, from) => {{
        const body = compactWithMap(raw.slice(from));
        if (!body.compact) return null;
        let best = null;
        for (const term of terms) {{
          const wanted = compactWithMap(term).compact;
          if (!wanted) continue;
          const idx = body.compact.indexOf(wanted);
          if (idx < 0) continue;
          const start = from + body.map[idx];
          const end = from + body.map[idx + wanted.length - 1] + 1;
          if (!best || start < best.start || (start === best.start && end > best.end)) {{
            best = {{ term, start, end }};
          }}
        }}
        return best;
      }};
      for (const node of nodes) {{
        if (!node.parentNode || !node.parentElement.getClientRects().length || getComputedStyle(node.parentElement).visibility === "hidden") continue;
        const raw = node.nodeValue || '';
        if (!raw.trim()) continue;
        const frag = document.createDocumentFragment();
        let cursor = 0;
        let matched = false;
        while (cursor < raw.length) {{
          const next = findNext(raw, cursor);
          if (!next || next.end <= cursor) break;
          if (next.start > cursor) frag.appendChild(document.createTextNode(raw.slice(cursor, next.start)));
          const mark = document.createElement('mark');
          mark.className = 'codex-highlight';
          mark.id = 'codex-highlight-' + markId++;
          mark.textContent = raw.slice(next.start, next.end);
          frag.appendChild(mark);
          termHits[next.term] += 1;
          cursor = next.end;
          matched = true;
        }}
        if (!matched) continue;
        if (cursor < raw.length) frag.appendChild(document.createTextNode(raw.slice(cursor)));
        node.parentNode.replaceChild(frag, node);
      }}
      window.__codexHighlightCount = markId;
      window.__codexTermHits = termHits;
      window.__codexMissingTerms = terms.filter((term) => !termHits[term]);
    }})();
    """
    if soup.body:
        soup.body.append(script)
    else:
        soup.append(script)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(str(soup), encoding="utf-8")
    return out_path


def screenshot_html(
    page,
    html_path: Path,
    out_path: Path,
    allow_unhighlighted: bool = False,
    allow_missing_terms: bool = False,
) -> dict:
    page.goto(html_path.as_uri(), wait_until="load")
    page.evaluate("async () => { if (document.fonts) await document.fonts.ready; }")
    page.wait_for_timeout(200)
    count = int(page.evaluate("window.__codexHighlightCount || 0"))
    missing_terms = page.evaluate("window.__codexMissingTerms || []")
    term_hits = page.evaluate("window.__codexTermHits || {}")
    if missing_terms and not (allow_missing_terms or allow_unhighlighted):
        raise ValueError(f"HTML terms not found in {html_path.name}: {missing_terms}")
    if count <= 0:
        if not allow_unhighlighted:
            raise ValueError(f"No HTML highlights were found for {html_path.name}")
        page.screenshot(path=str(out_path), full_page=False)
        return {"highlight_count": count, "missing_terms": missing_terms, "term_hits": term_hits}

    clip = page.evaluate(
        """() => {
          const marks = Array.from(document.querySelectorAll('mark.codex-highlight'));
          const blocks = marks.map(m => m.closest('table') || m.closest('p,li,blockquote,h1,h2,h3,h4') || m);
          const rects = blocks.map((m) => m.getBoundingClientRect()).filter((r) => r.width > 0 && r.height > 0);
          if (!rects.length) return null;
          const doc = document.documentElement;
          const body = document.body || doc;
          const pageWidth = Math.max(doc.scrollWidth, body.scrollWidth, doc.clientWidth);
          const pageHeight = Math.max(doc.scrollHeight, body.scrollHeight, doc.clientHeight);
          let left = 0;
          let right = pageWidth;
          let top = Math.min(...rects.map((r) => r.top + window.scrollY));
          let bottom = Math.max(...rects.map((r) => r.bottom + window.scrollY));
          const padX = 160, padY = 90;
          left = Math.max(0, left - padX);
          right = Math.min(pageWidth, right + padX);
          top = Math.max(0, top - padY);
          bottom = Math.min(pageHeight, bottom + padY);
          const minWidth = Math.min(1100, pageWidth);
          if (right - left < minWidth) {
            const center = (left + right) / 2;
            left = Math.max(0, center - minWidth / 2);
            right = Math.min(pageWidth, left + minWidth);
            left = Math.max(0, right - minWidth);
          }
          const minHeight = Math.min(240, pageHeight);
          if (bottom - top < minHeight) {
            const center = (top + bottom) / 2;
            top = Math.max(0, center - minHeight / 2);
            bottom = Math.min(pageHeight, top + minHeight);
            top = Math.max(0, bottom - minHeight);
          }
          return {x: Math.round(left), y: Math.round(top), width: Math.round(right-left), height: Math.round(bottom-top)};
        }"""
    )
    if not clip:
        if not allow_unhighlighted:
            raise ValueError(f"Cannot crop HTML screenshot because no highlight clip was found for {html_path.name}")
        page.screenshot(path=str(out_path), full_page=False)
        return {"highlight_count": count, "missing_terms": missing_terms, "term_hits": term_hits}

    if clip["height"] > 3500:
        raise ValueError("Evidence span is too tall; split sources or use a focused official PDF/table image")
    page.screenshot(path=str(out_path), clip=clip, full_page=True)
    return {"highlight_count": count, "missing_terms": missing_terms, "term_hits": term_hits}


def crop_image(path: Path, out_path: Path, clip: list[int] | None = None) -> None:
    with Image.open(path) as img:
        if clip:
            x, y, w, h = [int(v) for v in clip]
            if x < 0 or y < 0 or w <= 0 or h <= 0 or x+w > img.width or y+h > img.height:
                raise ValueError("Crop rectangle must be inside the source image")
            left = max(0, min(x, img.width - 1))
            top = max(0, min(y, img.height - 1))
            right = max(left + 1, min(left + w, img.width))
            bottom = max(top + 1, min(top + h, img.height))
            img.crop((left, top, right, bottom)).save(out_path)
        else:
            img.save(out_path)


def validate_source_meta(evidence_id: str, meta: dict, source_path: Path) -> None:
    language = str(meta.get("source_language", "")).lower()
    translated = meta.get("translation_zh", "")
    excerpt = meta.get("source_excerpt", "")
    if language.startswith("en") or language == "mixed":
        if not isinstance(translated, str) or not translated.strip():
            raise ValueError(f"{evidence_id}: English evidence requires translation_zh")
    if translated:
        if not isinstance(translated, str) or not isinstance(excerpt, str) or not excerpt.strip():
            raise ValueError(f"{evidence_id}: translation requires source_excerpt")
    suffix = source_path.suffix.lower()
    allow_unhighlighted = bool(meta.get("allow_unhighlighted"))
    pre_highlighted = bool(meta.get("pre_highlighted"))

    if suffix in {".html", ".htm"} and not meta.get("terms") and not allow_unhighlighted:
        raise ValueError(f"{evidence_id}: HTML source requires non-empty terms unless allow_unhighlighted=true")
    if suffix == ".pdf" and not meta.get("clip") and not allow_unhighlighted:
        raise ValueError(f"{evidence_id}: PDF source requires clip coordinates unless allow_unhighlighted=true")
    if suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        if not meta.get("clip") and not (pre_highlighted or allow_unhighlighted):
            raise ValueError(f"{evidence_id}: image source requires clip coordinates or pre_highlighted=true")


def render_pdf_source(source_path: Path, out_path: Path, page_index: int, clip: list[int] | None = None) -> None:
    try:
        import fitz  # type: ignore
    except ImportError as exc:
        raise RuntimeError("PDF sources require PyMuPDF/fitz. Install PyMuPDF or provide a pre-cropped image.") from exc

    with fitz.open(source_path) as doc:
        if page_index < 0 or page_index >= len(doc):
            raise ValueError(f"PDF page_index {page_index} out of range for {source_path} with {len(doc)} pages")
        page = doc[page_index]
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
    tmp = out_path.with_name(f"{out_path.stem}_rendered.png")
    pix.save(tmp)
    crop_image(tmp, out_path, clip)
    tmp.unlink(missing_ok=True)


def clean_source_label(meta: dict, source_path: Path) -> str:
    label = meta.get("label") or source_path.stem
    forbidden = ["网页位置", "：prepared", "（html", "(html"]
    lowered = str(label).lower()
    if any(term in lowered for term in forbidden):
        raise ValueError(f"Invalid source label; keep only the file name without type/location descriptions: {label}")
    return str(label)


def source_note(meta: dict, source_path: Path) -> str:
    label = clean_source_label(meta, source_path)
    page = str(meta.get("page", "")).strip()
    if page and page != "无页码":
        return f"资料来源：{label}，页码：{page}"
    return f"资料来源：{label}"


def make_source_assets(cfg: dict, dirs: dict[str, Path], base_dir: Path) -> dict[str, dict]:
    evidence_cfg: dict = cfg["evidence"]
    manifest: dict[str, dict] = {}
    browser_exe = os.environ.get("REPORT_BROWSER")
    html_items = []

    for evidence_id, meta in evidence_cfg.items():
        source = resolve_path(meta["source"], base_dir)
        validate_source_meta(evidence_id, meta, source)
        out = dirs["source"] / f"{evidence_id}.png"
        suffix = source.suffix.lower()
        if suffix in {".html", ".htm"}:
            html_items.append((evidence_id, meta, source, out))
        elif suffix == ".pdf":
            page_index = int(meta.get("page_index", 0))
            render_pdf_source(source, out, page_index, meta.get("clip"))
            manifest[evidence_id] = {
                "image": str(out),
                "note": source_note(meta, source),
                "calc": meta.get("calc", ""),
                "highlight_count": 1 if meta.get("clip") else 0,
                "allow_unhighlighted": bool(meta.get("allow_unhighlighted")),
            }
        elif suffix in {".png", ".jpg", ".jpeg", ".webp"}:
            crop_image(source, out, meta.get("clip"))
            manifest[evidence_id] = {
                "image": str(out),
                "note": source_note(meta, source),
                "calc": meta.get("calc", ""),
                "highlight_count": 1 if meta.get("clip") or meta.get("pre_highlighted") else 0,
                "allow_unhighlighted": bool(meta.get("allow_unhighlighted")),
            }
        else:
            raise ValueError(f"Unsupported source type for {source}")

    if html_items:
        viewport = cfg.get("source_viewport", {"width": 1400, "height": 1000})
        with sync_playwright() as p:
            browser = p.chromium.launch(headless=True, **({"executable_path": browser_exe} if browser_exe else {}))
            page = browser.new_page(
                viewport={"width": int(viewport.get("width", 1400)), "height": int(viewport.get("height", 1000))},
                device_scale_factor=1,
            )
            page.route("http://**/*", lambda route: route.abort())
            page.route("https://**/*", lambda route: route.abort())
            for evidence_id, meta, source, out in html_items:
                html = highlighted_html(source, dirs["html"] / f"{evidence_id}.html", meta.get("terms", []))
                shot = screenshot_html(
                    page,
                    html,
                    out,
                    allow_unhighlighted=bool(meta.get("allow_unhighlighted")),
                    allow_missing_terms=bool(meta.get("allow_missing_terms")),
                )
                manifest[evidence_id] = {
                    "image": str(out),
                    "note": source_note(meta, source),
                    "calc": meta.get("calc", ""),
                    "highlight_count": shot["highlight_count"],
                    "missing_terms": shot["missing_terms"],
                    "term_hits": shot["term_hits"],
                    "allow_unhighlighted": bool(meta.get("allow_unhighlighted")),
                }
            browser.close()

    for evidence_id, item in manifest.items():
        with Image.open(item["image"]) as img:
            item["width"] = img.width
            item["height"] = img.height
        meta = evidence_cfg[evidence_id]
        if meta.get("translation_zh"):
            translated = make_text_image(
                "", "中文译文（辅助翻译）", meta["translation_zh"],
                dirs["source"] / f"{evidence_id}_translation.png",
                include_sheet_title=False, width=900,
            )
            item["translation_image"] = translated["path"]
            item["translation_zh"] = meta["translation_zh"]
            item["source_excerpt"] = meta["source_excerpt"]
            item["source_language"] = meta.get("source_language", "")
    return manifest


def validate_config(cfg: dict, base_dir: Path, output_xlsx: Path, process_dir: Path) -> None:
    validate_output_paths(output_xlsx, process_dir, cfg)
    if output_xlsx.suffix.lower() != ".xlsx":
        raise ValueError("output_xlsx must end in .xlsx")
    if output_xlsx.exists():
        raise ValueError("Output already exists; choose a new version")
    if (process_dir / "visual_workpaper_manifest.json").exists():
        raise ValueError("Process directory belongs to an earlier output; choose a new version directory")
    if cfg.get("word_confirmed") is not True:
        auth = cfg.get("generation_authorization", {})
        if auth.get("mode") != "delegated_review" or auth.get("authorized") is not True or not str(auth.get("user_instruction", "")).strip() or not str(auth.get("scope", "")).strip():
            raise ValueError("Require genuine Word confirmation or explicit delegated-review authorization")
        review_path = resolve_path(auth.get("review_record", ""), base_dir)
        if not review_path.is_file():
            raise ValueError("Delegated review record is missing")
        review = json.loads(review_path.read_text(encoding="utf-8-sig"))
        if review.get("passed") is not True or review.get("word_sha256") != cfg.get("word_sha256") or not str(review.get("reviewer", "")).strip():
            raise ValueError("Delegated review must pass and identify the reviewer and exact Word hash")
    word_path = resolve_path(cfg["word_path"], base_dir)
    if not word_path.exists():
        raise FileNotFoundError(f"Word file not found: {word_path}")

    if hashlib.sha256(word_path.read_bytes()).hexdigest() != cfg.get("word_sha256"):
        raise ValueError("Word hash changed or missing; refresh mapping and the applicable confirmation or review record")
    paragraphs = extract_docx_paragraphs(word_path)
    sections = cfg.get("sections", [])
    if [sec.get("sheet") for sec in sections] != ["事件", "投资要点第一段", "投资要点第二段", "投资要点第三段"]:
        raise ValueError("Require the event and three investment sections in order")
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(word_path) as z:
        document = ET.fromstring(z.read("word/document.xml"))
    controls = {}
    for control in document.findall(".//w:sdt", ns):
        alias = control.find("w:sdtPr/w:alias", ns)
        content = control.find("w:sdtContent", ns)
        if alias is not None and content is not None:
            controls[alias.get("{" + ns["w"] + "}val")] = ["".join(t.text or "" for t in par.findall(".//w:t", ns)) for par in content.findall(".//w:p", ns)]
    if "事件公告" in controls and "摘要" in controls:
        expected = ["".join(controls["事件公告"])] + controls["摘要"][1:6:2]
        if [select_body(paragraphs, sec) for sec in sections] != expected:
            raise ValueError("Configured sections do not match the four final Word bodies")
    review_mode = any("review_inputs" in point or "review_result" in point for sec in sections for point in sec.get("points", []))
    evidence = cfg.get("evidence", {})
    if not evidence:
        raise ValueError("Config must include evidence")

    evidence_ids = set(evidence)
    if len({x.get("sheet") for x in sections}) != len(sections):
        raise ValueError("Duplicate sheet names")
    for name in evidence:
        if not re.fullmatch(r"[A-Za-z0-9_-]+", name):
            raise ValueError("Evidence IDs must use ASCII letters, digits, _ or -")
    for section in sections:
        sheet = section.get("sheet", "<missing sheet>")
        if not isinstance(sheet, str) or not sheet or len(sheet)>31 or re.search(r"[\\/*?:\[\]]",sheet):
            raise ValueError("Invalid Excel sheet name")
        body = select_body(paragraphs, section)
        overview = section.get("overview_evidence")
        if overview not in evidence_ids:
            raise ValueError(f"{sheet}: overview_evidence not found in evidence: {overview}")
        points = section.get("points") or []
        if not points:
            raise ValueError(f"{sheet}: must include at least one point")
        if all(point.get("left_mode", "point_text") == "point_text" for point in points) and "".join(point.get("text", "") for point in points) != body:
            raise ValueError(f"{sheet}: points must cover the complete body in order")
        for idx, point in enumerate(points, start=1):
            if review_mode and not all(isinstance(point.get(k), str) and point[k].strip() for k in ("review_inputs", "review_result")):
                raise ValueError("Summary-first layout requires review_inputs and review_result for every point")
            if not point.get("evidence_ids"):
                raise ValueError("Every point requires evidence_ids")
            selected = point.get("highlights", []) if point.get("left_mode") == "full_body_highlight" else [point.get("text", "")]
            if not selected or any(not text or text not in body for text in selected):
                raise ValueError("Point text is not in the confirmed Word paragraph")
            mode = point.get("left_mode", "point_text")
            if mode == "full_body_highlight":
                if not (section.get("body") or section.get("body_selector")):
                    raise ValueError(
                        f"{sheet} point {idx}: left_mode=full_body_highlight requires section body or body_selector"
                    )
                if not point.get("highlights"):
                    raise ValueError(f"{sheet} point {idx}: left_mode=full_body_highlight requires highlights")
            else:
                point_body_for_image(section, point)
            for ev_id in point.get("evidence_ids", []):
                if ev_id not in evidence_ids:
                    raise ValueError(f"{sheet} point {idx}: evidence_id not found: {ev_id}")

    for evidence_id, meta in evidence.items():
        source = resolve_path(meta["source"], base_dir)
        if not source.exists():
            raise FileNotFoundError(f"{evidence_id}: source file not found: {source}")
        validate_source_meta(evidence_id, meta, source)
        source_note(meta, source)


def scaled_dims(path: Path, max_width: int, max_height: int | None = None) -> tuple[int, int]:
    with Image.open(path) as img:
        w, h = img.size
    scale = max_width / w
    if max_height is not None:
        scale = min(scale, max_height / h)
    return max(1, int(w * scale)), max(1, int(h * scale))


def add_image(ws, image_path: Path, cell: str, max_width: int, max_height: int | None = None) -> tuple[int, int]:
    width, height = scaled_dims(image_path, max_width, max_height)
    img = XLImage(str(image_path))
    img.width = width
    img.height = height
    ws.add_image(img, cell)
    return width, height


def rows_for_px(px: int, note_lines: int = 0) -> int:
    return max(8, math.ceil((px + 42 + note_lines * 24) / 24))


def style_note_cell(cell, fill: str = "FFFFFF") -> None:
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.font = Font(name="Microsoft YaHei", size=10, color="333333")
    cell.fill = PatternFill("solid", fgColor=fill)
    side = Side(style="thin", color="D9D9D9")
    cell.border = Border(left=side, right=side, top=side, bottom=side)


def add_separator(ws, row: int, bilingual: bool = False) -> None:
    ws.row_dimensions[row].height = 12
    fill = PatternFill("solid", fgColor="FFF2CC")
    side = Side(style="thin", color="F1C232")
    for col in (("A", "B", "C", "D", "E") if bilingual else ("A", "B", "C", "D")):
        cell = ws[f"{col}{row}"]
        cell.fill = fill
        cell.border = Border(left=side, right=side, top=side, bottom=side)


def configure_sheet(ws, bilingual: bool = False) -> None:
    ws.sheet_view.showGridLines = False
    for col, width in {"A": 94, "B": 4, "C": 128, "D": 28}.items():
        ws.column_dimensions[col].width = width
    if bilingual:
        ws.column_dimensions["D"].width = 4
        ws.column_dimensions["E"].width = 94
    ws.sheet_format.defaultRowHeight = 18
    ws.freeze_panes = "A1"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.2
    ws.page_margins.right = 0.2
    ws.page_margins.top = 0.25
    ws.page_margins.bottom = 0.25


def point_body_for_image(section: dict, point: dict) -> str:
    mode = point.get("left_mode", "point_text")
    if mode == "full_body_highlight":
        return section["body"]
    if mode != "point_text":
        raise ValueError(f"Unsupported point left_mode: {mode}")
    if point.get("text"):
        return point["text"]
    highlights = point.get("highlights") or []
    if len(highlights) == 1:
        return highlights[0]
    raise ValueError(
        f"Point in sheet {section.get('sheet')} must provide text, "
        "or exactly one highlight, unless left_mode=full_body_highlight"
    )


def make_text_assets(sections: list[dict], dirs: dict[str, Path]) -> tuple[dict, dict]:
    para_manifest: dict[str, dict] = {}
    point_manifest: dict[str, list[dict]] = {}
    for section in sections:
        sheet = section["sheet"]
        para_manifest[sheet] = make_text_image(
            sheet,
            section.get("heading", sheet),
            section["body"],
            dirs["paragraph"] / f"{sheet}.png",
            include_sheet_title=True,
            width=900,
        )
        point_manifest[sheet] = []
        for idx, point in enumerate(section["points"], start=1):
            point_body = point_body_for_image(section, point)
            point_highlights = point.get("highlights", [])
            if point.get("left_mode", "point_text") == "point_text" and not point_highlights:
                point_highlights = [point_body]
            point_manifest[sheet].append(
                make_text_image(
                    sheet,
                    section.get("heading", sheet),
                    point_body,
                    dirs["point"] / f"{sheet}_{idx:02d}.png",
                    highlights=point_highlights,
                    include_sheet_title=(idx == 1),
                    width=900,
                )
            )
    return para_manifest, point_manifest


def create_legacy_workbook(
    output_xlsx: Path,
    sections: list[dict],
    para_manifest: dict,
    point_manifest: dict,
    source_manifest: dict,
) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    for section in sections:
        ws = wb.create_sheet(section["sheet"])
        used_ids = [section["overview_evidence"]] + [ev for point in section["points"] for ev in point["evidence_ids"]]
        bilingual = any(source_manifest[ev].get("translation_image") for ev in used_ids)
        configure_sheet(ws, bilingual)
        source_width = 900 if bilingual else 1180

        para_info = para_manifest[section["sheet"]]
        first_evidence = source_manifest[section["overview_evidence"]]
        _, para_h = add_image(ws, Path(para_info["path"]), "A1", max_width=660)
        note_rows = 1 + (1 if first_evidence.get("calc") else 0)
        ws["C1"] = first_evidence["note"]
        style_note_cell(ws["C1"])
        image_row = 2
        if first_evidence.get("calc"):
            ws["C2"] = first_evidence["calc"]
            style_note_cell(ws["C2"], fill="FFF2CC")
            image_row = 3
        _, ev_h = add_image(ws, Path(first_evidence["image"]), f"C{image_row}", max_width=source_width, max_height=520)
        if first_evidence.get("translation_image"):
            _, translated_h = add_image(ws, Path(first_evidence["translation_image"]), f"E{image_row}", max_width=660)
            ev_h = max(ev_h, translated_h)
        start_note_row = max(rows_for_px(para_h), rows_for_px(ev_h) + note_rows)
        current_row = start_note_row + 4

        for idx, point in enumerate(section["points"]):
            point_info = point_manifest[section["sheet"]][idx]
            _, left_h = add_image(ws, Path(point_info["path"]), f"A{current_row}", max_width=660)
            left_rows = rows_for_px(left_h)

            right_row = current_row
            for ev_id in point["evidence_ids"]:
                ev = source_manifest[ev_id]
                note_rows = 1
                ws[f"C{right_row}"] = ev["note"]
                style_note_cell(ws[f"C{right_row}"])
                if ev.get("calc"):
                    ws[f"C{right_row + 1}"] = ev["calc"]
                    style_note_cell(ws[f"C{right_row + 1}"], fill="FFF2CC")
                    note_rows += 1
                image_row = right_row + note_rows + 1
                _, height = add_image(ws, Path(ev["image"]), f"C{image_row}", max_width=source_width, max_height=520)
                if ev.get("translation_image"):
                    _, translated_h = add_image(ws, Path(ev["translation_image"]), f"E{image_row}", max_width=660)
                    height = max(height, translated_h)
                needed = note_rows + 1 + rows_for_px(height)
                right_row += needed + 2

            group_rows = max(left_rows, right_row - current_row)
            if idx < len(section["points"]) - 1:
                separator_row = current_row + group_rows
                add_separator(ws, separator_row, bilingual)
                current_row = separator_row + 3
            else:
                current_row += group_rows + 1

        last_column = "E" if bilingual else "D"
        ws.print_area = f"A1:{last_column}{max(current_row, 40)}"

    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)


def create_workbook(output_xlsx, sections, para_manifest, point_manifest, source_manifest):
    if not any("review_inputs" in p for s in sections for p in s["points"]):
        return create_legacy_workbook(output_xlsx, sections, para_manifest, point_manifest, source_manifest)
    wb = Workbook()
    wb.remove(wb.active)

    def text(ws, address, value, heading=False):
        cell = ws[address]
        cell.value = value
        cell.font = Font(name="Microsoft YaHei", size=18, bold=heading)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor="FFF2CC" if heading else "FFFFFF")

    def evidence_size(ev):
        _, height = scaled_dims(Path(ev["image"]), 900, 520)
        th = scaled_dims(Path(ev["translation_image"]), 660)[1] if ev.get("translation_image") else 0
        note_height = max(30, math.ceil(len(ev["note"]) / 36) * 28)
        return max(height, th), note_height, (30 + note_height) * 4 / 3 + max(height, th) + 72

    def reserve(ws, row, block):
        # ponytail: A3 layout uses a conservative pixel budget; verify the exported pages.
        if block > 1250:
            raise ValueError("Evidence group exceeds a print page; split the source/translation")
        start = getattr(ws, "_evidence_page_start", 1)
        used = sum((ws.row_dimensions[i].height or 18) * 4 / 3 for i in range(start, row))
        if used + block > 1250 and row > start:
            ws.row_breaks.append(Break(id=row - 1))
            ws._evidence_page_start = row

    def evidence(ws, row, key, label):
        ev = source_manifest[key]
        height, note_height, block = evidence_size(ev)
        reserve(ws, row, block)
        text(ws, f"C{row}", label + "｜" + key, True)
        text(ws, f"E{row}", "中文译文（辅助翻译）" if ev.get("translation_image") else "", True)
        ws.row_dimensions[row].height = 30
        text(ws, f"C{row + 1}", ev["note"])
        ws.row_dimensions[row + 1].height = note_height
        add_image(ws, Path(ev["image"]), f"C{row + 2}", 900, 520)
        if ev.get("translation_image"):
            add_image(ws, Path(ev["translation_image"]), f"E{row + 2}", 660)
        return row + 2 + math.ceil(height / 24) + 3

    for section in sections:
        ws = wb.create_sheet(section["sheet"])
        configure_sheet(ws, True)
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.sheet_view.zoomScale = 65
        _, h = add_image(ws, Path(para_manifest[section["sheet"]]["path"]), "A1", 660)
        row = max(evidence(ws, 1, section["overview_evidence"], "整段概览"), math.ceil(h / 24) + 3)
        add_separator(ws, row, True)
        row += 2
        for i, point in enumerate(section["points"], 1):
            values = [point_body_for_image(section, point), point["review_inputs"], point["review_result"]]
            lines = max(sum(max(1, math.ceil(sum(1 if ord(c) > 255 else .55 for c in line) / width)) for line in value.split("\n")) for value, width in zip(values, [26, 36, 26]))
            height = max(96, lines * 28 + 10)
            if height > 400:
                raise ValueError("Review summary is too long for one Excel row; split the sentence group")
            first_block = evidence_size(source_manifest[point["evidence_ids"][0]])[2]
            reserve(ws, row, (30 + height + 18) * 4 / 3 + first_block)
            for col, value in [("A", f"{i:02d}｜" + point.get("review_title", "逐句核对")), ("C", "核对数据（对应下方来源编号）"), ("E", "计算与结论（辅助核验）")]:
                text(ws, f"{col}{row}", value, True)
            ws.row_dimensions[row].height = 30
            for col, value in zip(["A", "C", "E"], values):
                text(ws, f"{col}{row + 1}", value)
            ws.row_dimensions[row + 1].height = height
            row += 3
            for n, key in enumerate(point["evidence_ids"], 1):
                row = evidence(ws, row, key, f"来源{n}")
            add_separator(ws, row, True)
            row += 2
        ws.print_area = f"A1:E{row - 1}"
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_xlsx)
    actual = load_workbook(output_xlsx)
    for section in sections:
        values = [cell.value for row in actual[section["sheet"]] for cell in row]
        for point in section["points"]:
            if not all(point[k] in values for k in ("review_inputs", "review_result")):
                raise ValueError("Exported review summary differs from configuration")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--config", required=True, help="Path to workpaper JSON config.")
    parser.add_argument("--validate-only", action="store_true", help="Validate config paths/references without generating Excel.")
    args = parser.parse_args()

    config_path = Path(args.config).resolve()
    cfg = load_config(config_path)
    base_dir = resolve_path(cfg.get("company_dir", "."), Path(cfg["_config_dir"]))
    output_xlsx = resolve_path(cfg["output_xlsx"], base_dir)
    process_dir = resolve_path(cfg.get("process_dir", "写作产出/_底稿过程文件"), base_dir)
    validate_config(cfg, base_dir, output_xlsx, process_dir)
    if args.validate_only:
        print(json.dumps({"validated": str(config_path)}, ensure_ascii=False, indent=2))
        return
    dirs = {
        "process": process_dir,
        "paragraph": process_dir / "paragraph_images",
        "point": process_dir / "point_images",
        "source": process_dir / "source_screenshots",
        "html": process_dir / "highlighted_html",
    }
    for path in dirs.values():
        path.mkdir(parents=True, exist_ok=True)

    word_path = resolve_path(cfg["word_path"], base_dir)
    paragraphs = extract_docx_paragraphs(word_path)
    sections = cfg["sections"]
    for section in sections:
        section["body"] = select_body(paragraphs, section)

    para_manifest, point_manifest = make_text_assets(sections, dirs)
    source_manifest = make_source_assets(cfg, dirs, base_dir)
    create_workbook(output_xlsx, sections, para_manifest, point_manifest, source_manifest)

    manifest = {
        "word_sha256": cfg["word_sha256"],
        "xlsx_sha256": hashlib.sha256(output_xlsx.read_bytes()).hexdigest(),
        "config_sha256": hashlib.sha256(config_path.read_bytes()).hexdigest(),
        "word_confirmed": cfg.get("word_confirmed") is True,
        "generation_authorization": cfg.get("generation_authorization"),
        "word": str(word_path),
        "output": str(output_xlsx),
        "sections": sections,
        "source_screenshots": source_manifest,
        "paragraph_images": para_manifest,
        "point_images": point_manifest,
    }
    (process_dir / "visual_workpaper_manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"output": str(output_xlsx)}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
