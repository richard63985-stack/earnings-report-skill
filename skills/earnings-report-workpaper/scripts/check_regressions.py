"""Synthetic regression check. Optional --browser checks off-screen HTML capture."""
from pathlib import Path
import argparse,copy,hashlib,json,os,subprocess,sys,tempfile,zipfile
import fitz
from PIL import Image
from openpyxl import load_workbook
import fill_report_template as fill
import create_visual_workpaper as paper
from make_template import create


def rejected(action, message):
    try:
        action()
    except (ValueError, FileNotFoundError) as error:
        assert message in str(error), str(error)
    else:
        raise AssertionError('Expected rejection: ' + message)


def run(base, browser=False):
    base.mkdir(parents=True,exist_ok=True)
    template=base/'template.docx';create(template)
    with zipfile.ZipFile(template) as z:xml=z.read('word/document.xml').decode()
    for alias,value in [('副标题','保留副标题'),('投资评级','已有评级'),('评级变动','已有变动')]:xml=fill.replace_sdt_single(xml,alias,value)
    fixed=base/'fixed.docx';fill.write_docx(template,fixed,xml)
    bodies=['虚构公司本季收入1.26亿元，同比增长26%。','订阅收入占比为60%。','AI产品付费客户达到24家。','毛利率为72%，同比提升4个百分点。']
    md='# 星帆测试公司：收入增长，客户拓展\n\n## 事件\n\n'+bodies[0]+'\n\n## 投资要点\n\n'
    for i in range(1,4):md+=f'▌虚构要点{i}\n\n{bodies[i]}\n\n'
    md+='## 风险提示\n\n全部为虚构测试数据。\n'
    draft=base/'draft.md';draft.write_text(md,encoding='utf-8')
    word=base/'report.docx';args=fill.build_parser().parse_args(['--template',str(fixed),'--markdown',str(draft),'--output',str(word),'--date','2031年07月01日'])
    assert fill.fill(args)==0
    with zipfile.ZipFile(word) as z:out=z.read('word/document.xml').decode()
    for alias in ['副标题','投资评级','评级变动']:
        a,b=fill.sdt_bounds_for_alias(xml,alias);c,d=fill.sdt_bounds_for_alias(out,alias);assert xml[a:b]==out[c:d]
    bad=base/'bad.md';bad.write_text(md.replace('▌虚构要点3\n\n'+bodies[3]+'\n\n','')+'\n## 投资建议\n\n不能用建议凑满三点。',encoding='utf-8')
    rejected(lambda:fill.parse_markdown(bad),'three nonempty')
    originals=['Revenue: Q2 2031 126, Q2 2030 100 (RMB million).','Subscription: 75.6; total revenue: 126 (RMB million).','Paying AI customers: 24 at June 30, 2031.','Gross margin: Q2 2031 72%, Q2 2030 68%.']
    inputs=['来源1：2031Q2收入126、2030Q2收入100，人民币百万元。','来源1：订阅75.6、总收入126，人民币百万元。','来源1：2031年6月30日AI付费客户24家。','来源1：本期72%、上年同期68%，均为毛利率。']
    results=['(126÷100−1)×100%=26%；126百万元=1.26亿元。','75.6÷126×100%=60%。','直接披露24家。','72%−68%=4个百分点。']
    ev={};sections=[]
    for i,name in enumerate(['事件','投资要点第一段','投资要点第二段','投资要点第三段']):
        source=base/f'source{i}.pdf';doc=fitz.open();page=doc.new_page();page.insert_text((35,40),'Fictional earnings release - synthetic data only');page.insert_text((35,65),originals[i]);doc.save(source);doc.close()
        ev[f'e{i}']={'source':source.name,'label':source.name,'page_index':0,'clip':[30,30,1100,160],'source_language':'en','source_excerpt':originals[i],'translation_zh':inputs[i].replace('来源1：','')}
        sections.append({'sheet':name,'body':bodies[i],'overview_evidence':f'e{i}','points':[{'text':bodies[i],'evidence_ids':[f'e{i}'],'review_inputs':inputs[i],'review_result':results[i]}]})
    cfg={'company_dir':'.','word_path':word.name,'word_sha256':hashlib.sha256(word.read_bytes()).hexdigest(),'word_confirmed':True,'output_xlsx':'workpaper.xlsx','process_dir':'process-v1','evidence':ev,'sections':sections}
    def validate(c):paper.validate_config(c,base,base/c['output_xlsx'],base/c['process_dir'])
    validate(cfg)
    for edit,message in [(lambda c:c.update(word_confirmed=False),'explicit delegated-review'),(lambda c:c.update(word_sha256='0'*64),'hash changed'),(lambda c:c['sections'][0]['points'][0].pop('review_result'),'requires review_inputs'),(lambda c:c['sections'][0]['points'][0].update(text='虚构公司本季收入1.26亿元，'),'complete body'),(lambda c:c['sections'].reverse(),'sections in order')]:
        changed=copy.deepcopy(cfg);edit(changed);rejected(lambda:validate(changed),message)
    delegated=copy.deepcopy(cfg);delegated['word_confirmed']=False
    delegated['generation_authorization']={'mode':'delegated_review','authorized':True,'user_instruction':'虚构测试授权主控复核后生成','scope':'仅此虚构样例','review_record':'review.json'}
    review=base/'review.json';review.write_text(json.dumps({'passed':True,'reviewer':'synthetic reviewer','word_sha256':cfg['word_sha256']}),encoding='utf-8');validate(delegated)
    review.write_text(json.dumps({'passed':True,'reviewer':'synthetic reviewer','word_sha256':'0'*64}),encoding='utf-8');rejected(lambda:validate(delegated),'exact Word hash')
    review.write_text(json.dumps({'passed':True,'reviewer':'synthetic reviewer','word_sha256':cfg['word_sha256']}),encoding='utf-8')
    config=base/'workpaper.json';config.write_text(json.dumps(delegated,ensure_ascii=False),encoding='utf-8')
    subprocess.run([sys.executable,str(Path(paper.__file__)),'--config',str(config)],check=True,stdout=subprocess.DEVNULL)
    wb=load_workbook(base/'workpaper.xlsx');assert wb.sheetnames==[s['sheet']for s in sections]
    with zipfile.ZipFile(base/'workpaper.xlsx')as z:embedded={hashlib.sha256(z.read(n)).hexdigest()for n in z.namelist()if n.startswith('xl/media/')}
    manifest=json.loads((base/'process-v1/visual_workpaper_manifest.json').read_text(encoding='utf-8'))
    assert manifest['word_confirmed'] is False and manifest['word_sha256']==cfg['word_sha256']
    assert manifest['xlsx_sha256']==hashlib.sha256((base/'workpaper.xlsx').read_bytes()).hexdigest()
    for s in sections:
        ws=wb[s['sheet']];values=[c.value for row in ws for c in row];point=s['points'][0]
        assert all(point[k]in values for k in ['text','review_inputs','review_result'])
        for key in point['evidence_ids']:
            for kind in ['image','translation_image']:assert hashlib.sha256(Path(manifest['source_screenshots'][key][kind]).read_bytes()).hexdigest()in embedded
    rejected(lambda:validate(cfg),'Output already exists')
    changed=copy.deepcopy(cfg);changed['output_xlsx']='workpaper-v2.xlsx';rejected(lambda:validate(changed),'earlier output')
    # Legacy input still generates its existing layout.
    legacy=copy.deepcopy(cfg);legacy.update(output_xlsx='legacy.xlsx',process_dir='legacy-process')
    for s in legacy['sections']:
        for p in s['points']:p.pop('review_inputs');p.pop('review_result')
    (base/'legacy.json').write_text(json.dumps(legacy,ensure_ascii=False),encoding='utf-8')
    subprocess.run([sys.executable,str(Path(paper.__file__)),'--config',str(base/'legacy.json')],check=True,stdout=subprocess.DEVNULL)
    assert len(load_workbook(base/'legacy.xlsx').sheetnames)==4
    if browser:
        source=base/'long.html';source.write_text('<html><body style="margin:0"><div style="height:4200px;background:#0000ff"></div><p style="margin:0;height:140px;background:#00ff00">Revenue 126. Complete paragraph ending here.</p><div style="height:1000px;background:#ff0000"></div></body></html>')
        prepared=base/'highlight.html';paper.highlighted_html(source,prepared,['Revenue 126'])
        with paper.sync_playwright()as pw:
            b=pw.chromium.launch(headless=True,**({'executable_path':os.environ['REPORT_BROWSER']}if os.environ.get('REPORT_BROWSER')else {}));page=b.new_page(viewport={'width':1200,'height':800},device_scale_factor=1)
            shot=base/'long-crop.png';paper.screenshot_html(page,prepared,shot)
            with Image.open(shot)as im:
                pixels=list(im.convert('RGB').getdata());assert sum(g>200 and r<50 and b<50 for r,g,b in pixels)>len(pixels)*.2
            missing=base/'missing.html';paper.highlighted_html(source,missing,['does not exist']);rejected(lambda:paper.screenshot_html(page,missing,base/'missing.png'),'terms not found')
            b.close()
    print(json.dumps({'synthetic_regressions':'passed','browser_checked':browser,'output':str(base)},ensure_ascii=False))


if __name__=='__main__':
    for stream in (sys.stdout, sys.stderr):
        stream.reconfigure(encoding='utf-8')
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output-dir');p.add_argument('--browser',action='store_true');args=p.parse_args()
    if args.output_dir:run(Path(args.output_dir).resolve(),args.browser)
    else:
        with tempfile.TemporaryDirectory(prefix='earnings-skill-check-')as d:run(Path(d),args.browser)
