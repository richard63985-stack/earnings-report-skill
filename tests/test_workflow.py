import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

SCRIPTS=Path(__file__).resolve().parents[1]/'skills/earnings-report-workpaper/scripts'
sys.path.insert(0,str(SCRIPTS))
import create_visual_workpaper as wp
from make_demo import create_demo
from freeze_word import freeze
from fill_report_template import ps_quote, export_pdf_with_word
from openpyxl import load_workbook
import fitz


class Workflow(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp=tempfile.TemporaryDirectory(prefix='report 中文 space ')
        cls.base=create_demo(Path(cls.tmp.name)/'case')
        cls.cfg=json.loads((cls.base/'workpaper.json').read_text(encoding='utf-8'))

    @classmethod
    def tearDownClass(cls): cls.tmp.cleanup()

    def config(self): return json.loads(json.dumps(self.cfg))

    def validate(self,cfg):
        wp.validate_config(cfg,self.base,self.base/'unused.xlsx',self.base/'process')

    def test_word_and_hash(self):
        data=freeze(self.base/'report.docx')
        self.assertFalse(data['word_confirmed'])
        self.assertEqual(data['word_sha256'],self.cfg['word_sha256'])
        for section in self.cfg['sections']:self.assertIn(section['body'],data['paragraphs'])
        self.validate(self.config())

    def test_wrong_hash_rejected(self):
        cfg=self.config();cfg['word_sha256']='0'*64
        with self.assertRaisesRegex(ValueError,'hash'):self.validate(cfg)

    def test_without_confirmation_rejected(self):
        cfg=self.config();cfg['word_confirmed']=False
        with self.assertRaisesRegex(ValueError,'confirm'):self.validate(cfg)

    def test_stale_body_rejected(self):
        cfg=self.config();cfg['sections'][0]['body']='旧正文。'
        with self.assertRaisesRegex(ValueError,'differs'):self.validate(cfg)

    def test_missing_evidence_rejected(self):
        cfg=self.config();cfg['sections'][0]['points'][0]['evidence_ids']=[]
        with self.assertRaisesRegex(ValueError,'requires evidence'):self.validate(cfg)

    def test_fabricated_point_rejected(self):
        cfg=self.config();cfg['sections'][0]['points'][0]['text']='新增的无出处句子。'
        with self.assertRaisesRegex(ValueError,'not in'):self.validate(cfg)

    def test_path_traversal_rejected(self):
        cfg=self.config();cfg['evidence']['../escape']=cfg['evidence']['source_0']
        with self.assertRaisesRegex(ValueError,'Evidence IDs'):self.validate(cfg)

    def test_submit_folder_rejected(self):
        with self.assertRaisesRegex(ValueError,'提交材料'):
            wp.validate_output_paths(self.base/'提交材料'/'report.xlsx',self.base/'process',{})

    def test_highlight_missing_rejected(self):
        with self.assertRaisesRegex(ValueError,'Highlight not found'):
            wp.highlight_ranges('原文',['不存在'],strict=True)

    def test_pdf_crop(self):
        path=self.base/'source.pdf';out=self.base/'crop.png'
        doc=fitz.open();page=doc.new_page(width=400,height=300);page.insert_text((30,50),'Fictional revenue: 100');doc.save(path);doc.close()
        wp.render_pdf_source(path,out,0,[20,40,500,120])
        from PIL import Image
        with Image.open(out) as img:self.assertEqual(img.size,(500,120))
        with self.assertRaisesRegex(ValueError,'out of range'):wp.render_pdf_source(path,out,99,[0,0,1,1])

    def test_powershell_quote(self):
        self.assertEqual(ps_quote("a'b $x"),"'a''b $x'")
        if sys.platform!='win32':self.assertFalse(export_pdf_with_word(Path('a'),Path('b')))

    def test_end_to_end_browser_excel(self):
        subprocess.run([sys.executable,str(SCRIPTS/'create_visual_workpaper.py'),'--config',str(self.base/'workpaper.json')],check=True)
        path=self.base/'workpaper.xlsx'
        with zipfile.ZipFile(path) as z:self.assertIsNone(z.testzip())
        wb=load_workbook(path)
        self.assertEqual(wb.sheetnames,[s['sheet'] for s in self.cfg['sections']])
        for ws in wb:
            self.assertEqual(len(ws._images),4)
            self.assertFalse(any(c.data_type=='e' for row in ws for c in row))
        wb.close()
        manifest=json.loads((self.base/'process/visual_workpaper_manifest.json').read_text(encoding='utf-8'))
        self.assertEqual(len(manifest['source_screenshots']),4)
        for item in manifest['source_screenshots'].values():self.assertGreater(item['highlight_count'],0)


if __name__=='__main__':unittest.main()
